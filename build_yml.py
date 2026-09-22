#!/usr/bin/env python3
"""Собрать Yandex YML из книги Google Sheets.

B24U рисует карточку товара только из тега <picture> в YML.
Колонка image в CSV в карточку не попадает — отсюда «нет картинок».

Источник: xlsx-выгрузка таблицы
https://docs.google.com/spreadsheets/d/1pDfZ5RIsfbeIUrhdIoeRJ9efb2DyxqUT18CzOqoOpvw
"""

from __future__ import annotations

import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

SOURCE_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1pDfZ5RIsfbeIUrhdIoeRJ9efb2DyxqUT18CzOqoOpvw"
)

SHOPS: dict[str, tuple[str, str, str]] = {
    # sheet → (slug, shop name, shop url)
    "Yard76": ("yard76", "Yard76", "https://yard76.ru"),
    "Zeleniy-magazin": ("zelenyi-magazin", "Зелёный магазин", "https://zelenyi-magazin.ru"),
    "Южный": ("yuzhniy", "Южный", "https://www.uzhniy.ru"),
    "httpsoracdecor.ru": ("oracdecor", "Orac Decor", "https://oracdecor.ru"),
    "Смазки": ("smazki", "Смазка.ру", "https://smazka.ru.com"),
    "Нипол": ("nipol", "НИПОЛ", "https://nikart74.ru"),
}

PARAM_LABELS = {
    "code": "Код",
    "weight": "Вес",
    "width": "Ширина",
    "height": "Высота",
    "length": "Длина",
    "country": "Страна",
    "import_country": "Страна ввоза",
    "special_prices": "Спеццены",
    "params": "Параметры",
    "latin_name": "Латинское название",
    "characteristics": "Характеристики",
    "parent_id": "Родительский товар",
    "source_category": "Раздел",
    "material": "Материал",
    "dimensions": "Размеры",
    "length_cm": "Длина, см",
    "height_cm": "Высота, см",
    "width_cm": "Ширина, см",
    "price_unit": "Единица цены",
    "price_per_m": "Цена за метр",
    "tags": "Метки",
    "stock": "Остаток",
}

SKIP_COLUMNS = {
    "id",
    "sku",
    "article",
    "query_sku",
    "product_id",
    "name",
    "title",
    "price",
    "old_price",
    "currency",
    "url",
    "description",
    "available",
    "category",
    "brand",
    "vendor",
    "image",
    "status",
    "exact_article_match",
    "flex_url",
    "magento_sku",
}

TRUE_WORDS = {"true", "1", "да", "yes", "в наличии", "есть"}
FALSE_WORDS = {"false", "0", "нет", "no", "недоступен", "нет в наличии"}


def cell(row: tuple, index: dict[str, int], key: str):
    pos = index.get(key)
    if pos is None or pos >= len(row):
        return None
    value = row[pos]
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    return value


def as_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def as_number(value) -> str | None:
    """Число для <price>/<oldprice>. Дату и «По запросу» не превращаем в цену."""
    if value is None or isinstance(value, (bool, datetime, date)):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != value:
            return None
        if float(value) < 0:
            return None
        if float(value).is_integer():
            return str(int(value))
        return f"{float(value):.2f}".rstrip("0").rstrip(".")
    text = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not re.fullmatch(r"\d+(?:\.\d+)?", text):
        return None
    if "." in text:
        return text.rstrip("0").rstrip(".")
    return text


def as_available(value) -> str:
    if value is None:
        return "true"
    if isinstance(value, bool):
        return "true" if value else "false"
    text = str(value).strip().lower()
    if text in TRUE_WORDS:
        return "true"
    if text in FALSE_WORDS:
        return "false"
    return "true"


def pictures(value) -> list[str]:
    if value is None:
        return []
    found = re.findall(r"https?://[^\s,;|]+", str(value))
    out: list[str] = []
    for url in found:
        url = url.rstrip(").,]")
        if url not in out:
            out.append(url)
    return out


def add(parent: ET.Element, tag: str, value: str | None) -> None:
    if value is None or value == "":
        return
    node = ET.SubElement(parent, tag)
    node.text = value


def offer_identity(row, index: dict[str, int], fallback: int) -> str:
    for key in ("id", "sku", "article", "product_id", "query_sku"):
        value = cell(row, index, key)
        if value is not None:
            return as_text(value)
    return str(fallback)


def vendor_code(row, index: dict[str, int]) -> str | None:
    for key in ("sku", "article", "query_sku"):
        value = cell(row, index, key)
        if value is not None:
            return as_text(value)
    return None


class Categories:
    def __init__(self) -> None:
        self._ids: dict[str, str] = {}
        self.nodes: list[tuple[str, str | None, str]] = []

    def ensure(self, path: str | None) -> str:
        parts = [part.strip() for part in re.split(r"\s*/\s*", path or "") if part.strip()]
        if not parts:
            parts = ["Каталог"]
        parent: str | None = None
        acc: list[str] = []
        for part in parts:
            acc.append(part)
            key = "/".join(acc)
            if key not in self._ids:
                cid = str(len(self._ids) + 1)
                self._ids[key] = cid
                self.nodes.append((cid, parent, part))
            parent = self._ids[key]
        return parent or "1"


def build_sheet(ws, sheet_name: str) -> tuple[str, dict]:
    slug, shop_name, shop_url = SHOPS[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h else "" for h in rows[0]]
    index = {name: pos for pos, name in enumerate(headers) if name}

    categories = Categories()
    offers = ET.Element("offers")
    seen: dict[str, int] = {}
    stats = {
        "slug": slug,
        "shop": shop_name,
        "rows": 0,
        "offers": 0,
        "with_picture": 0,
        "without_picture": 0,
        "without_price": 0,
        "skipped": 0,
    }

    for row_no, row in enumerate(rows[1:], start=2):
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        stats["rows"] += 1
        status = cell(row, index, "status")
        if status is not None and str(status).strip().lower() not in {"", "ok"}:
            stats["skipped"] += 1
            continue
        name = cell(row, index, "name") or cell(row, index, "title")
        url = cell(row, index, "url")
        if name is None or url is None:
            stats["skipped"] += 1
            continue

        raw_id = offer_identity(row, index, row_no)
        seen[raw_id] = seen.get(raw_id, 0) + 1
        offer_id = raw_id if seen[raw_id] == 1 else f"{raw_id}-{seen[raw_id]}"

        # Порядок как у рабочего фида prosps: карточка B24U = name + url + picture
        # + vendor + vendorCode. <param> в эту карточку не входит — фото остаётся ссылкой.
        offer = ET.SubElement(
            offers,
            "offer",
            {"id": offer_id, "available": as_available(cell(row, index, "available"))},
        )
        add(offer, "name", as_text(name))
        add(offer, "url", as_text(url))

        pics = pictures(cell(row, index, "image"))
        for pic in pics:
            add(offer, "picture", pic)
        if pics:
            stats["with_picture"] += 1
        else:
            stats["without_picture"] += 1

        add(offer, "currencyId", "RUB")
        add(
            offer,
            "categoryId",
            categories.ensure(
                as_text(cell(row, index, "category")) if cell(row, index, "category") else None
            ),
        )
        brand = cell(row, index, "brand") or cell(row, index, "vendor")
        add(offer, "vendor", as_text(brand) if brand is not None else shop_name)
        add(offer, "vendorCode", vendor_code(row, index) or offer_id)

        price = as_number(cell(row, index, "price"))
        old_price = as_number(cell(row, index, "old_price"))
        if price is None:
            stats["without_price"] += 1
        else:
            add(offer, "price", price)
            if old_price is not None and float(old_price) > float(price):
                add(offer, "oldprice", old_price)

        stock_number = as_number(cell(row, index, "stock"))
        add(offer, "quantity", stock_number or "1")

        description = cell(row, index, "description")
        desc_text = re.sub(r"\s+", " ", as_text(description)) if description is not None else ""
        extras: list[str] = []
        if price is None:
            raw_price = cell(row, index, "price")
            if isinstance(raw_price, str) and raw_price.strip():
                extras.append(f"Цена: {raw_price.strip()}")
        for key, label in PARAM_LABELS.items():
            if key in SKIP_COLUMNS or key not in index or key == "stock":
                continue
            value = cell(row, index, key)
            if value is None or isinstance(value, (datetime, date)):
                continue
            text = as_text(value)
            if text and text not in desc_text:
                extras.append(f"{label}: {text}")
        if extras:
            tail = ". ".join(extras)
            desc_text = f"{desc_text} {tail}".strip() if desc_text else tail
        add(offer, "description", desc_text or None)

        stats["offers"] += 1

    now = datetime.now().astimezone().isoformat(timespec="seconds")
    root = ET.Element("yml_catalog", {"date": now})
    shop = ET.SubElement(root, "shop")
    add(shop, "name", shop_name)
    add(shop, "company", shop_name)
    add(shop, "url", shop_url)
    add(shop, "platform", "BSM/Yandex/Market")
    currencies = ET.SubElement(shop, "currencies")
    ET.SubElement(currencies, "currency", {"id": "RUB", "rate": "1"})
    categories_el = ET.SubElement(shop, "categories")
    for cid, parent, title in categories.nodes:
        attrib = {"id": cid}
        if parent:
            attrib["parentId"] = parent
        node = ET.SubElement(categories_el, "category", attrib)
        node.text = title
    shop.append(offers)

    ET.indent(root, space="  ")
    xml = ET.tostring(root, encoding="unicode", xml_declaration=False)
    body = '<?xml version="1.0" encoding="utf-8"?>\n' + xml + "\n"
    stats["categories"] = len(categories.nodes)
    return body, stats


def write_index(out_dir: Path, all_stats: list[dict]) -> None:
    rows = []
    for item in all_stats:
        url = f"{item['slug']}/feed.xml"
        rows.append(
            "<tr>"
            f"<td>{item['shop']}</td>"
            f"<td><a href=\"{url}\">{url}</a></td>"
            f"<td>{item['offers']}</td>"
            f"<td>{item['with_picture']}</td>"
            f"<td>{item['without_picture']}</td>"
            f"<td>{item['without_price']}</td>"
            "</tr>"
        )
    html = f"""<!doctype html>
<html lang="ru">
<meta charset="utf-8">
<title>YML-фиды каталогов</title>
<style>
  body {{ font: 16px/1.45 -apple-system, BlinkMacSystemFont, sans-serif; margin: 2rem auto; max-width: 60rem; padding: 0 1rem; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border-bottom: 1px solid #ddd; text-align: left; padding: 0.4rem 0.6rem; }}
  code {{ font-size: 0.95em; }}
</style>
<h1>YML-фиды для B24U</h1>
<p>Каждый лист таблицы — отдельный фид в формате Яндекс Маркета. Карточка виджета берёт фото из <code>&lt;picture&gt;</code>, валюта только <code>RUB</code>.</p>
<p>Источник: <a href="{SOURCE_URL}">Google Sheets</a>.</p>
<table>
  <tr><th>Магазин</th><th>Фид</th><th>Офферы</th><th>С фото</th><th>Без фото</th><th>Без числовой цены</th></tr>
  {''.join(rows)}
</table>
"""
    (out_dir / "index.html").write_text(html, encoding="utf-8")


def main() -> None:
    source = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/feeds.xlsx")
    out_dir = Path(sys.argv[2] if len(sys.argv) > 2 else Path(__file__).resolve().parent)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    all_stats = []
    for sheet_name in workbook.sheetnames:
        if sheet_name not in SHOPS:
            print(f"skip unknown sheet {sheet_name}")
            continue
        body, stats = build_sheet(workbook[sheet_name], sheet_name)
        target = out_dir / stats["slug"] / "feed.xml"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(body, encoding="utf-8")
        all_stats.append(stats)
        print(
            f"{stats['slug']}: offers={stats['offers']} pictures={stats['with_picture']} "
            f"no_picture={stats['without_picture']} no_price={stats['without_price']} "
            f"skipped={stats['skipped']} → {target}"
        )
    write_index(out_dir, all_stats)


if __name__ == "__main__":
    main()
